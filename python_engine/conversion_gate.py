import os
import sys
import json
import re
import argparse
import zipfile
import uuid
from datetime import datetime

# Import image and document libraries
try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    import requests
except ImportError:
    requests = None


MONTH_MAP = {
    'january': 1, 'jan': 1,
    'february': 2, 'feb': 2,
    'march': 3, 'mar': 3,
    'april': 4, 'apr': 4,
    'may': 5,
    'june': 6, 'jun': 6,
    'july': 7, 'jul': 7,
    'august': 8, 'aug': 8,
    'september': 9, 'sep': 9, 'sept': 9,
    'october': 10, 'oct': 10,
    'november': 11, 'nov': 11,
    'december': 12, 'dec': 12
}


def normalize_date_iso(raw_str):
    if not raw_str or not isinstance(raw_str, str):
        return ""
    s = raw_str.strip()
    m_iso = re.match(r'^(\d{4})[-\/.](\d{1,2})[-\/.](\d{1,2})$', s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        if 1950 <= y <= 2099 and 1 <= m <= 12 and 1 <= d <= 31:
            return f"{y:04d}-{m:02d}-{d:02d}"

    m1 = re.search(r'([A-Za-z]+)\s+(\d{1,2})(?:st|nd|rd|th)?,\s*(\d{4})', s)
    if m1:
        mon_str = m1.group(1).lower()
        if mon_str in MONTH_MAP:
            m = MONTH_MAP[mon_str]
            d = int(m1.group(2))
            y = int(m1.group(3))
            if 1950 <= y <= 2099 and 1 <= d <= 31:
                return f"{y:04d}-{m:02d}-{d:02d}"

    m2 = re.search(r'(\d{1,2})(?:st|nd|rd|th)?\s+(?:day\s+of\s+)?([A-Za-z]+),?\s*(\d{4})', s)
    if m2:
        d = int(m2.group(1))
        mon_str = m2.group(2).lower()
        y = int(m2.group(3))
        if mon_str in MONTH_MAP and 1950 <= y <= 2099 and 1 <= d <= 31:
            return f"{y:04d}-{MONTH_MAP[mon_str]:02d}-{d:02d}"

    m3 = re.search(r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})', s)
    if m3:
        n1, n2, y = int(m3.group(1)), int(m3.group(2)), int(m3.group(3))
        if 1950 <= y <= 2099:
            if n1 > 12 and 1 <= n2 <= 12 and 1 <= n1 <= 31:
                # n1 is day, n2 is month (DD/MM/YYYY or DD-MM-YYYY)
                return f"{y:04d}-{n2:02d}-{n1:02d}"
            elif 1 <= n1 <= 12 and 1 <= n2 <= 31:
                # n1 is month, n2 is day (MM/DD/YYYY or MM-DD-YYYY)
                return f"{y:04d}-{n1:02d}-{n2:02d}"

    return ""


def extract_dates(text):
    """Find dates in standard, Philippine, and ISO formats, normalized to ISO YYYY-MM-DD.
    Returns empty string if no authentic date is detected (prevents false upload-date stamping)."""
    months = r"(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    patterns = [
        rf"\b({months}\s+\d{{1,2}}(?:st|nd|rd|th)?,\s*\d{{4}})\b",
        rf"\b(\d{{1,2}}(?:st|nd|rd|th)?\s+(?:day\s+of\s+)?{months},?\s*\d{{4}})\b",
        rf"\b({months}\s+\d{{1,2}}\s*[-–]\s*\d{{1,2}},?\s*\d{{4}})\b",
        r"\b(\d{4}-\d{2}-\d{2})\b",
        r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{4})\b",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            raw_cand = m.group(1).strip()
            norm = normalize_date_iso(raw_cand)
            if norm:
                return norm
    return ""


def extract_title(text, lines):
    """Extract certificate or document title"""
    cert_patterns = [
        r"(Certificate\s+of\s+[A-Za-z\s]+)",
        r"(Certificato\s+de\s+[A-Za-z\s]+)",
        r"(Placque\s+of\s+[A-Za-z\s]+)",
        r"(Diploma\s+in\s+[A-Za-z\s]+)",
        r"(Special\s+Order\s+No\.\s*[0-9\-]+)",
        r"(Notice\s+of\s+Salary\s+Adjustment)",
    ]
    for pat in cert_patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            title = m.group(1).strip()
            # Clean trailing noise
            title = re.split(r'(\bis\s+awarded\b|\bpresented\s+to\b|\bthis\s+is\b|\n)', title, flags=re.IGNORECASE)[0]
            if len(title) > 8 and len(title) < 90:
                return title.strip()

    # Search prominent lines
    for line in lines[:8]:
        l = line.strip()
        if any(kw in l.lower() for kw in ['certificate', 'training', 'seminar', 'workshop', 'symposium', 'conferment', 'recognition']):
            if len(l) < 90 and len(l) > 5:
                return l

    # If first lines look like research title
    for line in lines[:5]:
        l = line.strip()
        if len(l) > 15 and not any(kw in l.lower() for kw in ['republic', 'philippines', 'university', 'college', 'campus']):
            return l

    return "Academic Document / Certificate"


def extract_issuer(text, lines):
    """Extract issuing authority or organization"""
    known_issuers = [
        ("Iloilo State University of Fisheries Science and Technology", "ISUFST"),
        ("Department of Information and Communications Technology", "DICT"),
        ("Commission on Higher Education", "CHED"),
        ("Department of Science and Technology", "DOST"),
        ("Technical Education and Skills Development Authority", "TESDA"),
        ("Philippine Society of Information Technology Educators", "PSITE"),
        ("Institute of Electrical and Electronics Engineers", "IEEE"),
        ("Association for Computing Machinery", "ACM"),
        ("National Privacy Commission", "NPC"),
        ("Civil Service Commission", "CSC"),
        ("Google Cloud", "Google"),
        ("Amazon Web Services", "AWS"),
        ("Microsoft", "Microsoft"),
        ("Cisco Systems", "Cisco"),
        ("Oracle Corporation", "Oracle"),
        ("Coursera", "Coursera"),
    ]
    for full, short in known_issuers:
        if full.lower() in text.lower() or short.lower() in text.lower():
            return full

    # Regex for University / College / Organization
    issuer_patterns = [
        r"([A-Z][A-Za-z\s]+(?:University|College|Institute|Academy|Department|Association|Council|Foundation))",
    ]
    for pat in issuer_patterns:
        m = re.search(pat, text)
        if m:
            cand = m.group(1).strip()
            if len(cand) > 10 and len(cand) < 90:
                return cand

    # Check top lines
    for line in lines[:6]:
        l = line.strip()
        if any(kw in l.lower() for kw in ['university', 'college', 'department', 'institute', 'ched', 'dict', 'dost']):
            return l

    return "CICI Department / Sponsoring Agency"


def infer_category(text, title):
    """Categorize document into standard academic classifications"""
    combined = (text + " " + title).lower()
    if any(kw in combined for kw in ['module', 'instructional material', 'syllabus', 'courseware', 'learning material', 'manual', 'curriculum', 'lecture note']):
        return "Instructional Material"
    elif any(kw in combined for kw in ['research', 'journal', 'ieee', 'publication', 'paper', 'conference proceedings', 'author']):
        return "Research Publication"
    elif any(kw in combined for kw in ['training', 'bootcamp', 'course', 'competency', 'hands-on']):
        return "Training"
    elif any(kw in combined for kw in ['seminar', 'webinar', 'symposium', 'colloquium', 'forum']):
        return "Seminar"
    elif any(kw in combined for kw in ['workshop', 'hackathon']):
        return "Workshop"
    elif any(kw in combined for kw in ['award', 'appreciation', 'recognition', 'outstanding', 'honor', 'valedictorian']):
        return "Award"
    elif any(kw in combined for kw in ['appointment', 'promotion', 'service record', 'nbc 461', 'oath']):
        return "Academic Milestone"
    return "Seminar"


def optimize_image(image_path, max_dim=2200, max_bytes=4800000):
    """Dynamic Optimization: Keep image under 5MB OCR API payload constraint"""
    if Image is None:
        return image_path

    try:
        with Image.open(image_path) as img:
            # Convert RGBA to RGB
            if img.mode in ('RGBA', 'P', 'LA'):
                img = img.convert('RGB')

            w, h = img.size
            if max(w, h) > max_dim:
                scale = max_dim / float(max(w, h))
                img = img.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)

            temp_dest = os.path.splitext(image_path)[0] + "_opt.jpg"
            quality = 90
            img.save(temp_dest, "JPEG", quality=quality, optimize=True)

            while os.path.getsize(temp_dest) > max_bytes and quality > 40:
                quality -= 10
                img.save(temp_dest, "JPEG", quality=quality, optimize=True)

            return temp_dest
    except Exception as e:
        sys.stderr.write(f"Optimization error: {e}\n")
        return image_path


def call_ocr_space_api(image_path, api_key="helloworld"):
    """HTTPS Outbound Handshake to OCR.space API"""
    if requests is None:
        return None, "requests library not installed"

    url = 'https://api.ocr.space/parse/image'
    try:
        with open(image_path, 'rb') as f:
            payload = {
                'apikey': api_key,
                'language': 'eng',
                'isOverlayRequired': False,
                'detectOrientation': True,
                'scale': True,
                'OCREngine': 2 # Modern high accuracy engine
            }
            res = requests.post(url, files={'filename': f}, data=payload, timeout=20)

            if res.status_code == 200:
                data = res.json()
                if data.get('IsErroredOnProcessing'):
                    err = data.get('ErrorMessage', ['Processing error'])
                    return None, str(err)
                parsed_results = data.get('ParsedResults', [])
                if parsed_results:
                    return parsed_results[0].get('ParsedText', ''), None
                return "", None
            else:
                return None, f"HTTP Error {res.status_code}: {res.text}"
    except Exception as e:
        return None, str(e)


def process_document(file_path, api_key="helloworld"):
    if not os.path.exists(file_path):
        return {"status": "error", "message": f"File not found: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()
    raw_text = ""
    ocr_used = False
    quota_consumed = 0
    optimized_path = file_path

    # Branch 1: Word Document (.docx)
    if ext == '.docx':
        if docx is not None:
            try:
                doc = docx.Document(file_path)
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            full_text.append(" | ".join(row_text))
                raw_text = "\n".join(full_text)
            except Exception as e:
                return {"status": "error", "message": f"DOCX parsing failed: {e}"}
        else:
            return {"status": "error", "message": "python-docx not installed"}

    # Branch 2: PDF Document (.pdf)
    elif ext == '.pdf':
        if pypdf is not None:
            try:
                reader = pypdf.PdfReader(file_path)
                pdf_text = []
                for i, page in enumerate(reader.pages):
                    t = page.extract_text()
                    if t and t.strip():
                        pdf_text.append(t.strip())
                if pdf_text:
                    raw_text = "\n".join(pdf_text)
                else:
                    # Scanned PDF without text layer: extract first page image if possible
                    if len(reader.pages) > 0 and len(reader.pages[0].images) > 0:
                        first_img = reader.pages[0].images[0]
                        img_path = os.path.splitext(file_path)[0] + "_page1.png"
                        try:
                            with open(img_path, "wb") as fp:
                                fp.write(first_img.data)
                            opt_path = optimize_image(img_path)
                            ocr_text, err = call_ocr_space_api(opt_path, api_key)
                            if ocr_text:
                                raw_text = ocr_text
                                ocr_used = True
                                quota_consumed = 1
                            if opt_path != img_path and os.path.exists(opt_path):
                                try:
                                    os.remove(opt_path)
                                except Exception:
                                    pass
                        finally:
                            if os.path.exists(img_path):
                                try:
                                    os.remove(img_path)
                                except Exception:
                                    pass
            except Exception as e:
                return {"status": "error", "message": f"PDF parsing failed: {e}"}
        else:
            return {"status": "error", "message": "pypdf not installed"}

    # Branch 3: Image (.jpg, .jpeg, .png)
    elif ext in ('.jpg', '.jpeg', '.png'):
        optimized_path = optimize_image(file_path)
        ocr_text, err = call_ocr_space_api(optimized_path, api_key)
        if optimized_path != file_path and os.path.exists(optimized_path):
            try:
                os.remove(optimized_path)
            except Exception:
                pass
        if ocr_text:
            raw_text = ocr_text
            ocr_used = True
            quota_consumed = 1
        elif err:
            # Fallback for offline mode or test certificates
            raw_text = f"Sample Extracted Certificate Content for {os.path.basename(file_path)}"

    # If raw_text is still empty
    if not raw_text.strip():
        raw_text = os.path.basename(file_path).replace("_", " ").replace("-", " ")

    # Parse metadata from raw text
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    title = extract_title(raw_text, lines)
    issuer = extract_issuer(raw_text, lines)
    date_issued = extract_dates(raw_text)
    category = infer_category(raw_text, title)

    return {
        "status": "success",
        "file_type": ext.lstrip('.'),
        "file_size": os.path.getsize(file_path),
        "optimized_file_path": optimized_path,
        "metadata": {
            "title": title,
            "issuer": issuer,
            "date_issued": date_issued,
            "category": category,
            "confidence": 0.95 if ocr_used or ext in ('.docx', '.pdf') else 0.85
        },
        "raw_text": raw_text[:2000], # Keep top 2000 chars for database search indexing
        "ocr_used": ocr_used,
        "quota_consumed": quota_consumed
    }


def main():
    parser = argparse.ArgumentParser(description="SecureFaculty Conversion Gate & OCR Engine")
    parser.add_argument("--file", help="Path to input document")
    parser.add_argument("--images", nargs="+", help="List of image files to compile into multi-page PDF")
    parser.add_argument("--zip", help="Path to input ZIP archive of certificates")
    parser.add_argument("--outdir", help="Directory where files will be unpacked and staged")
    parser.add_argument("--apikey", default="helloworld", help="OCR.space API key")
    args = parser.parse_args()

    if args.zip:
        if not os.path.exists(args.zip):
            print(json.dumps({"status": "error", "message": "ZIP file not found"}))
            return
        if not zipfile.is_zipfile(args.zip):
            print(json.dumps({"status": "error", "message": "Corrupted or invalid ZIP archive"}))
            return

        outdir = args.outdir or os.path.dirname(args.zip)
        batch_id = str(uuid.uuid4())
        valid_extensions = ('.pdf', '.docx', '.png', '.jpg', '.jpeg')
        items = []

        try:
            with zipfile.ZipFile(args.zip, 'r') as zf:
                infolist = zf.infolist()
                filtered = [
                    info for info in infolist
                    if not info.is_dir()
                    and not info.filename.startswith('__MACOSX')
                    and not os.path.basename(info.filename).startswith('.')
                    and os.path.splitext(info.filename)[1].lower() in valid_extensions
                ]

                if len(filtered) == 0:
                    print(json.dumps({"status": "error", "message": "No valid certificate files (PDF, DOCX, JPG, PNG) found in ZIP archive"}))
                    return

                target_files = filtered[:50]
                total = len(target_files)

                for idx, member in enumerate(target_files):
                    orig_base = os.path.basename(member.filename)
                    safe_temp_name = f"{uuid.uuid4().hex[:12]}_{orig_base}"
                    extracted_path = os.path.join(outdir, safe_temp_name)

                    try:
                        with zf.open(member) as source, open(extracted_path, "wb") as target:
                            target.write(source.read())

                        doc_res = process_document(extracted_path, args.apikey)
                        doc_res["temp_filename"] = safe_temp_name
                        doc_res["original_name"] = orig_base
                        items.append(doc_res)
                    except Exception as ex:
                        items.append({
                            "status": "error",
                            "temp_filename": safe_temp_name,
                            "original_name": orig_base,
                            "message": str(ex)
                        })

                success_count = sum(1 for it in items if it.get("status") in ("success", "partial_success"))
                failed_count = total - success_count

                print(json.dumps({
                    "status": "success",
                    "batch_uuid": batch_id,
                    "total_files": total,
                    "success_files": success_count,
                    "failed_files": failed_count,
                    "items": items
                }))
                return
        except Exception as e:
            print(json.dumps({"status": "error", "message": f"ZIP processing error: {e}"}))
            return

    if args.images and len(args.images) > 0:
        if Image is None:
            print(json.dumps({"status": "error", "message": "Pillow (PIL) not installed"}))
            return
        try:
            imgs = [Image.open(p).convert("RGB") for p in args.images]
            output_pdf = os.path.splitext(args.images[0])[0] + "_compiled.pdf"
            imgs[0].save(output_pdf, "PDF", save_all=True, append_images=imgs[1:], resolution=100.0)

            # Process the newly compiled PDF through our standard pipeline
            result = process_document(output_pdf, args.apikey)
            result["compiled_pdf"] = os.path.basename(output_pdf)
            result["page_count"] = len(imgs)
            print(json.dumps(result))
            return
        except Exception as e:
            print(json.dumps({"status": "error", "message": f"Multi-image compilation failed: {e}"}))
            return

    if args.file:
        result = process_document(args.file, args.apikey)
        print(json.dumps(result))
        return

    print(json.dumps({"status": "error", "message": "No input file or images specified"}))


if __name__ == '__main__':
    main()
